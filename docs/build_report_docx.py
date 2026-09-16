"""
docs.build_report_docx
======================

PURPOSE
-------
Render ``docs/PROJECT-REPORT.md`` as a Word document.

This machine has no pandoc and no LibreOffice, so the conversion is done here
with python-docx. The parser handles exactly the Markdown subset the report
uses and nothing more:

    # / ## / ### / ####     headings
    - item                  bullet list
    1. item                 numbered list
    > quote                 block quote
    | a | b |               pipe table (second row is the --- separator)
    ```                     fenced code block
    ---                     horizontal rule (rendered as a spacer)
    **bold**  *italic*  `code`   inline runs

Deliberately NOT supported: nested lists, images, links, HTML. The report does
not use them, and a general Markdown parser is a much larger thing than this
document needs.

USAGE
-----
    python docs/build_report_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SRC = Path(__file__).parent / "PROJECT-REPORT.md"
OUT = Path(__file__).parent / "AI-Data-Science-Workbench-Project-Report.docx"

INK = RGBColor(0x1A, 0x1F, 0x2B)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x5B, 0x64, 0x74)

# **bold** | *italic* | `code`  — one alternation, so a single pass splits the
# line into literal text and styled fragments without re-scanning.
INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")


def add_runs(paragraph, text: str) -> None:
    """Write `text` into `paragraph`, honouring inline bold/italic/code."""
    for fragment in INLINE.split(text):
        if not fragment:
            continue
        run = paragraph.add_run()
        if fragment.startswith("**") and fragment.endswith("**"):
            run.text = fragment[2:-2]
            run.font.bold = True
        elif fragment.startswith("`") and fragment.endswith("`"):
            run.text = fragment[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = ACCENT
        elif fragment.startswith("*") and fragment.endswith("*"):
            run.text = fragment[1:-1]
            run.font.italic = True
        else:
            run.text = fragment


def style_document(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12.5),
                       ("Heading 4", 11.5)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        style.paragraph_format.space_before = Pt(16 if size > 14 else 12)
        style.paragraph_format.space_after = Pt(6)


def add_table(doc, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            add_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if r == 0:
                    run.font.bold = True
    doc.add_paragraph()


def add_code(doc, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = INK


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def convert(src: Path, out: Path) -> Path:
    doc = Document()
    style_document(doc)

    lines = src.read_text(encoding="utf-8").splitlines()
    seen_title = False
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Fenced code block -------------------------------------------------
        if stripped.startswith("```"):
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
            i += 1
            continue

        # Pipe table --------------------------------------------------------
        if stripped.startswith("|") and i + 1 < len(lines) \
                and set(lines[i + 1].strip()) <= set("|-: "):
            rows = [split_row(stripped)]
            i += 2  # skip the --- separator row
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # Horizontal rule ---------------------------------------------------
        if stripped in ("---", "***", "___"):
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Headings ----------------------------------------------------------
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            # The document's first H1 becomes the Title style. A fresh
            # Document() has NO paragraphs, so this cannot be detected by
            # inspecting doc.paragraphs[0] — hence the explicit flag.
            if level == 1 and not seen_title:
                heading = doc.add_heading(level=0)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                seen_title = True
            else:
                heading = doc.add_heading(level=min(level, 4))
            add_runs(heading, text)
            for run in heading.runs:
                run.font.color.rgb = ACCENT
            i += 1
            continue

        # Block quote -------------------------------------------------------
        if stripped.startswith(">"):
            quoted = [stripped.lstrip("> ").rstrip()]
            i += 1
            while i < len(lines) and lines[i].strip().startswith(">"):
                quoted.append(lines[i].strip().lstrip("> ").rstrip())
                i += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(10)
            add_runs(paragraph, " ".join(quoted))
            for run in paragraph.runs:
                run.font.italic = True
                run.font.color.rgb = MUTED
            continue

        # Lists -------------------------------------------------------------
        if re.match(r"^[-*]\s+", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_runs(paragraph, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        # Blank line --------------------------------------------------------
        if not stripped:
            i += 1
            continue

        # Paragraph — join the soft-wrapped lines the Markdown source uses.
        buffer = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", ">", "|", "```", "---"))
                    or re.match(r"^([-*]|\d+\.)\s+", nxt)):
                break
            buffer.append(nxt)
            i += 1
        add_runs(doc.add_paragraph(), " ".join(buffer))

    doc.save(out)
    return out


if __name__ == "__main__":
    path = convert(SRC, OUT)
    document = Document(path)
    print(f"OK  {path}")
    print(f"    {len(document.paragraphs)} paragraphs, "
          f"{len(document.tables)} tables")
